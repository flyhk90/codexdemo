from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


OUT = Path("Serena_Codex_讲解.pptx")

WIDE_W = Inches(13.333)
WIDE_H = Inches(7.5)

BG = RGBColor(248, 250, 252)
DARK = RGBColor(20, 29, 45)
TEXT = RGBColor(45, 55, 72)
MUTED = RGBColor(101, 116, 139)
BLUE = RGBColor(37, 99, 235)
TEAL = RGBColor(13, 148, 136)
GREEN = RGBColor(22, 163, 74)
AMBER = RGBColor(217, 119, 6)
RED = RGBColor(220, 38, 38)
LINE = RGBColor(226, 232, 240)
WHITE = RGBColor(255, 255, 255)
CODE_BG = RGBColor(15, 23, 42)
PURPLE = RGBColor(124, 58, 237)


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color


def add_bg(slide, color=BG):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, WIDE_W, WIDE_H)
    set_fill(rect, color)
    rect.line.fill.background()
    rect.z_order = 0


def add_text(slide, text, x, y, w, h, size=24, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.65, 0.35, 7.8, 0.55, size=28, color=DARK, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.68, 0.9, 8.8, 0.35, size=12, color=MUTED)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.68), Inches(1.25), Inches(1.25), Inches(0.045))
    set_fill(line, BLUE)


def add_footer(slide, n):
    add_text(slide, f"{n:02d}", 12.35, 7.03, 0.35, 0.2, size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, title, body, accent=BLUE):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(card, WHITE)
    card.line.color.rgb = LINE
    card.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    set_fill(bar, accent)
    add_text(slide, title, x + 0.25, y + 0.18, w - 0.45, 0.32, size=15, color=DARK, bold=True)
    add_text(slide, body, x + 0.25, y + 0.6, w - 0.45, h - 0.75, size=10.5, color=TEXT)


def add_pill(slide, text, x, y, w, color):
    pill = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.34))
    set_fill(pill, color)
    add_text(slide, text, x, y + 0.06, w, 0.16, size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER)


def add_code(slide, code, x, y, w, h):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(box, CODE_BG)
    add_text(slide, code, x + 0.22, y + 0.22, w - 0.44, h - 0.44, size=12, color=RGBColor(226, 232, 240))


def add_arrow(slide, x1, y1, x2, y2, color=MUTED):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(2)
    line.line.end_arrowhead = True
    return line


def bullet_box(slide, items, x, y, w, h, size=14, color=TEXT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.level = 0
        p.space_after = Pt(8)
    return box


def slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, DARK)
    add_text(slide, "Serena", 0.75, 0.92, 4.8, 0.72, size=42, color=WHITE, bold=True)
    add_text(slide, "让 Codex 拥有 IDE 级代码理解能力", 0.82, 1.75, 7.2, 0.45, size=22, color=RGBColor(203, 213, 225))
    add_text(slide, "围绕本地项目的语义查找、引用分析、符号级编辑与项目记忆", 0.84, 2.32, 8.1, 0.35, size=13, color=RGBColor(148, 163, 184))
    add_code(
        slide,
        'serena.exe start-mcp-server --context codex\n'
        '--enable-web-dashboard false',
        0.85,
        4.65,
        5.45,
        1.15,
    )
    for x, y, w, label, color in [
        (7.6, 1.15, 2.6, "find_symbol", BLUE),
        (9.9, 2.25, 2.2, "references", TEAL),
        (7.25, 3.45, 2.5, "memory", GREEN),
        (10.1, 4.75, 2.05, "editing", AMBER),
    ]:
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.72))
        set_fill(shape, color)
        add_text(slide, label, x, y + 0.21, w, 0.2, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 1)


def slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "一句话定位", "Serena 是 Codex 连接本地代码库的语义操作层")
    add_card(slide, 0.72, 1.75, 3.8, 3.35, "不是搜索引擎", "它不负责查互联网文档，也不是替代 Context7。它关注你的本地项目：文件、符号、引用和修改点。", BLUE)
    add_card(slide, 4.85, 1.75, 3.8, 3.35, "更像 IDE 后端", "它提供类似跳转定义、查找引用、文件概览、按函数/类编辑的能力，让 AI 少做纯文本猜测。", TEAL)
    add_card(slide, 8.98, 1.75, 3.55, 3.35, "可沉淀项目知识", "通过 onboarding 和 memory 记录项目结构、测试命令、约定和长期上下文。", GREEN)
    add_footer(slide, 2)


def slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "它解决的问题", "大项目里，文本搜索能找到线索，但语义工具能降低误判")
    items = [
        ("定位慢", "函数名常见、目录复杂、同名概念多，全文搜索需要大量人工排除。", BLUE),
        ("引用不准", "字符串匹配无法区分定义、调用、注释、测试数据和旧代码。", RED),
        ("修改风险高", "直接替换文本容易改错范围，尤其是类方法、重载、嵌套函数。", AMBER),
        ("上下文易丢", "每次新任务都重新摸索项目结构、命令和约定，成本高。", PURPLE),
    ]
    y = 1.55
    for title, body, color in items:
        add_card(slide, 1.0, y, 11.25, 0.98, title, body, color)
        y += 1.18
    add_footer(slide, 3)


def slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Serena 在 Codex 中的位置", "MCP server 把本地语义能力暴露成工具")
    labels = [
        ("用户需求", 0.9, 2.8, BLUE),
        ("Codex", 3.1, 2.8, DARK),
        ("Serena MCP", 5.35, 2.8, TEAL),
        ("语言服务 / 索引", 7.85, 2.8, GREEN),
        ("本地代码库", 10.35, 2.8, AMBER),
    ]
    for label, x, y, color in labels:
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.65), Inches(0.86))
        set_fill(shape, color)
        add_text(slide, label, x, y + 0.28, 1.65, 0.2, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    for x1, x2 in [(2.55, 3.1), (4.75, 5.35), (7.0, 7.85), (9.5, 10.35)]:
        add_arrow(slide, x1, 3.23, x2, 3.23)
    add_code(
        slide,
        '[mcp_servers.serena]\n'
        'type = "stdio"\n'
        'command = "C:/Users/.../serena.exe"\n'
        'args = ["start-mcp-server", "--context", "codex"]',
        1.25,
        4.75,
        10.7,
        1.25,
    )
    add_footer(slide, 4)


def slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "核心工具矩阵", "你本机 Serena 暴露出来的主要能力")
    rows = [
        ("项目与目录", "activate_project, list_dir, find_file", "选定项目，快速定位文件"),
        ("符号理解", "get_symbols_overview, find_symbol", "看清文件结构，按类/函数查找"),
        ("引用分析", "find_referencing_symbols", "确认谁在调用某个符号"),
        ("代码修改", "replace_symbol_body, insert_before_symbol, insert_after_symbol", "按符号边界进行更精确编辑"),
        ("项目记忆", "onboarding, write_memory, read_memory", "保存项目约定和长期上下文"),
    ]
    x0, y0 = 0.72, 1.55
    widths = [2.05, 5.1, 4.7]
    headers = ["类别", "工具", "用途"]
    colors = [DARK, DARK, DARK]
    for i, h in enumerate(headers):
        add_pill(slide, h, x0 + sum(widths[:i]) + i * 0.12, y0, widths[i], colors[i])
    y = y0 + 0.52
    for idx, row in enumerate(rows):
        fill = WHITE if idx % 2 == 0 else RGBColor(241, 245, 249)
        for i, text in enumerate(row):
            rect = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                Inches(x0 + sum(widths[:i]) + i * 0.12),
                Inches(y),
                Inches(widths[i]),
                Inches(0.72),
            )
            set_fill(rect, fill)
            rect.line.color.rgb = LINE
            add_text(slide, text, x0 + sum(widths[:i]) + i * 0.12 + 0.12, y + 0.15, widths[i] - 0.24, 0.35, size=9.5, color=TEXT)
        y += 0.78
    add_footer(slide, 5)


def slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "和普通搜索的区别", "rg 很快，Serena 更懂代码结构")
    add_text(slide, "没有 Serena", 1.0, 1.55, 4.3, 0.35, size=18, color=DARK, bold=True)
    add_text(slide, "有 Serena", 7.35, 1.55, 4.3, 0.35, size=18, color=DARK, bold=True)
    left = [
        'rg "login\\("',
        "读取候选文件",
        "人工判断定义/调用/注释",
        "用测试和类型检查兜底",
    ]
    right = [
        "find_symbol(login)",
        "find_referencing_symbols",
        "按符号边界替换函数体",
        "结合测试验证行为",
    ]
    bullet_box(slide, left, 1.05, 2.2, 4.2, 2.9, size=15)
    bullet_box(slide, right, 7.4, 2.2, 4.2, 2.9, size=15)
    add_arrow(slide, 5.55, 3.35, 7.0, 3.35, BLUE)
    add_text(slide, "从文本匹配升级到语义定位", 5.18, 3.75, 2.25, 0.36, size=11, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 6)


def slide_7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "和 Context7 的区别", "一个看外部文档，一个看本地代码")
    add_card(slide, 0.9, 1.65, 5.35, 3.9, "Context7", "用途：查询库/框架的当前文档和代码示例\n\n典型问题：Next.js middleware 怎么配置？Supabase Auth 邮箱登录怎么写？\n\n关键词：官方文档、版本、API 示例", BLUE)
    add_card(slide, 7.1, 1.65, 5.35, 3.9, "Serena", "用途：理解和修改你的本地项目\n\n典型问题：这个项目里的 login 在哪？谁调用了它？怎么只替换这个函数体？\n\n关键词：代码库、符号、引用、项目记忆", TEAL)
    add_footer(slide, 7)


def slide_8(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "典型工作流", "从理解项目到完成修改")
    steps = [
        ("1", "激活项目", "activate_project"),
        ("2", "项目概览", "onboarding / list_dir"),
        ("3", "定位符号", "find_symbol"),
        ("4", "查找引用", "find_referencing_symbols"),
        ("5", "精确编辑", "replace_symbol_body / insert_after_symbol"),
        ("6", "验证结果", "测试、类型检查、运行应用"),
    ]
    y = 1.7
    for n, title, tool in steps:
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(1.0), Inches(y), Inches(0.52), Inches(0.52))
        set_fill(circle, BLUE if int(n) % 2 else TEAL)
        add_text(slide, n, 1.0, y + 0.14, 0.52, 0.12, size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, title, 1.75, y + 0.02, 2.6, 0.25, size=15, color=DARK, bold=True)
        add_text(slide, tool, 4.1, y + 0.04, 6.7, 0.24, size=12, color=MUTED)
        y += 0.78
    add_footer(slide, 8)


def slide_9(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "实战例子", "需求：找出 login 方法并评估修改影响")
    add_code(
        slide,
        "用户：找一下 login 方法在哪些地方被调用，改登录失败提示\n\n"
        "Serena 路线：\n"
        "1. find_symbol: 定位 login 定义\n"
        "2. find_referencing_symbols: 找调用方\n"
        "3. replace_symbol_body: 只改目标函数\n"
        "4. 运行测试或类型检查",
        0.85,
        1.65,
        5.95,
        3.95,
    )
    add_card(slide, 7.25, 1.75, 4.95, 1.1, "收益 1", "减少误改：不会把注释、文案或无关同名函数当成调用点。", GREEN)
    add_card(slide, 7.25, 3.1, 4.95, 1.1, "收益 2", "上下文更完整：先看谁调用，再决定修改是否影响接口契约。", BLUE)
    add_card(slide, 7.25, 4.45, 4.95, 1.1, "收益 3", "编辑范围更清晰：围绕符号边界操作，代码审查更容易。", TEAL)
    add_footer(slide, 9)


def slide_10(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "使用建议", "把 Serena 当成语义导航仪，而不是万能自动驾驶")
    add_card(slide, 0.95, 1.65, 3.65, 3.45, "适合优先用", "跨文件重构\n查找函数/类引用\n理解陌生项目\n按函数或类做局部修改\n沉淀项目命令和约定", GREEN)
    add_card(slide, 4.85, 1.65, 3.65, 3.45, "仍要配合", "rg 全文搜索\n测试和类型检查\n代码审查\n框架文档查询\n运行时日志", BLUE)
    add_card(slide, 8.75, 1.65, 3.65, 3.45, "注意边界", "动态调用和反射可能难以完整识别\n生成代码/超大仓库需要耐心\n未 onboarding 时上下文较少", AMBER)
    add_footer(slide, 10)


def slide_11(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, DARK)
    add_text(slide, "总结", 0.85, 0.85, 2.2, 0.55, size=32, color=WHITE, bold=True)
    bullet_box(
        slide,
        [
            "Serena 让 Codex 更像在使用 IDE：能按符号理解、查引用、做精确编辑。",
            "它不替代 rg、测试和文档查询，而是把本地代码理解这一步做得更稳。",
            "最佳组合：Context7 查外部库文档，Serena 理解本地项目，Codex 负责推理、实现和验证。",
        ],
        1.05,
        2.05,
        10.7,
        2.3,
        size=18,
        color=RGBColor(226, 232, 240),
    )
    add_code(slide, "Context7 = 外部文档\nSerena   = 本地代码语义\nCodex    = 协作执行与验证", 1.05, 5.15, 5.1, 1.1)
    add_footer(slide, 11)


def main():
    prs = Presentation()
    prs.slide_width = WIDE_W
    prs.slide_height = WIDE_H
    for fn in [
        slide_1,
        slide_2,
        slide_3,
        slide_4,
        slide_5,
        slide_6,
        slide_7,
        slide_8,
        slide_9,
        slide_10,
        slide_11,
    ]:
        fn(prs)
    prs.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
